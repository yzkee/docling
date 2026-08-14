from collections.abc import Iterable
from pathlib import Path
from types import SimpleNamespace

import pytest
from docling_core.types.doc import (
    ContentLayer,
    GroupItem,
    NodeItem,
    PictureClassificationLabel,
    PictureItem,
    TextItem,
)

from docling.backend.docx.drawingml.utils import get_libreoffice_cmd
from docling.backend.mspowerpoint_backend import MsPowerpointDocumentBackend
from docling.datamodel.backend_options import MsPowerpointBackendOptions
from docling.datamodel.base_models import InputFormat, ItemAndImageEnrichmentElement
from docling.datamodel.document import ConversionResult, DoclingDocument, InputDocument
from docling.datamodel.pipeline_options import ConvertPipelineOptions
from docling.document_converter import DocumentConverter, PowerpointFormatOption
from docling.models.base_model import BaseItemAndImageEnrichmentModel
from docling.pipeline.simple_pipeline import SimplePipeline

from .test_data_gen_flag import GEN_TEST_DATA
from .verify_utils import verify_document, verify_export

GENERATE = GEN_TEST_DATA

CHART_PPTX = Path("./tests/data/pptx/sources/pptx_chart.pptx")


class _PictureEnrichmentModel(BaseItemAndImageEnrichmentModel):
    images_scale = 1.0

    def is_processable(self, doc: DoclingDocument, element: NodeItem) -> bool:
        return isinstance(element, PictureItem)

    def __call__(
        self,
        doc: DoclingDocument,
        element_batch: Iterable[ItemAndImageEnrichmentElement],
    ) -> Iterable[NodeItem]:
        for element in element_batch:
            yield element.item


class _ChartEnrichmentPipeline(SimplePipeline):
    def __init__(self, pipeline_options: ConvertPipelineOptions) -> None:
        super().__init__(pipeline_options)
        self.enrichment_pipe = [_PictureEnrichmentModel()]


@pytest.fixture(scope="module")
def libreoffice_available() -> bool:
    """Return True when a working LibreOffice installation is detected."""
    try:
        return get_libreoffice_cmd(raise_if_unavailable=True) is not None
    except Exception:
        return False


def get_pptx_paths():
    # Define the directory you want to search
    directory = Path("./tests/data/pptx/sources/")

    # List all PPTX files in the directory and its subdirectories
    pptx_files = sorted(directory.rglob("*.pptx"))
    return pptx_files


def get_converter():
    from docling.document_converter import DocumentConverter

    converter = DocumentConverter(allowed_formats=[InputFormat.PPTX])

    return converter


def convert_with_pptx_backend(pptx_path: Path) -> DoclingDocument:
    in_doc = InputDocument(
        path_or_stream=pptx_path,
        format=InputFormat.PPTX,
        backend=MsPowerpointDocumentBackend,
    )

    assert in_doc.valid
    return in_doc._backend.convert()


def test_e2e_pptx_conversions():
    pptx_paths = get_pptx_paths()
    converter = get_converter()

    for pptx_path in pptx_paths:
        # print(f"converting {pptx_path}")

        gt_path = pptx_path.parent.parent / "groundtruth" / pptx_path.name

        conv_result: ConversionResult = converter.convert(pptx_path)

        doc: DoclingDocument = conv_result.document

        included_content_layers = (
            set(ContentLayer) if gt_path.stem in "powerpoint_comments" else None
        )
        pred_md: str = doc.export_to_markdown(
            compact_tables=True,
            included_content_layers=included_content_layers,
        )
        assert verify_export(
            pred_md,
            str(gt_path) + ".md",
            GENERATE,
        ), "export to md"

        pred_itxt: str = doc._export_to_indented_text(
            max_text_len=70, explicit_tables=False
        )
        assert verify_export(pred_itxt, str(gt_path) + ".itxt", GENERATE), (
            "export to indented-text"
        )

        assert verify_document(doc, str(gt_path) + ".json", GENERATE, fuzzy=True), (
            "document document"
        )


def test_comments_extraction() -> None:
    """Test comprehensive comment extraction including metadata, authors, and slide distribution."""

    converter = get_converter()
    path = Path("./tests/data/pptx/sources/powerpoint_comments.pptx")
    doc: DoclingDocument = converter.convert(path).document

    assert doc.num_pages() == 3, f"Expected 3 slides, got {doc.num_pages()}"

    # Comment groups: 4 total (2 on slide 1, 0 on slide 2, 2 on slide 3)
    comment_groups = [
        g
        for g in doc.groups
        if isinstance(g, GroupItem) and g.name.startswith("comment-")
    ]
    assert len(comment_groups) == 4, (
        f"Expected 4 comment groups, got {len(comment_groups)}"
    )

    assert all(g.content_layer == ContentLayer.NOTES for g in comment_groups), (
        "All comment groups should be in NOTES content layer"
    )

    slide1_comments = [g for g in comment_groups if "slide1" in g.name]
    slide2_comments = [g for g in comment_groups if "slide2" in g.name]
    slide3_comments = [g for g in comment_groups if "slide3" in g.name]
    assert len(slide1_comments) == 2, (
        f"Expected 2 comments on slide 1, got {len(slide1_comments)}"
    )
    assert len(slide2_comments) == 0, (
        f"Expected 0 comments on slide 2, got {len(slide2_comments)}"
    )
    assert len(slide3_comments) == 2, (
        f"Expected 2 comments on slide 3, got {len(slide3_comments)}"
    )

    comment_texts = [
        t.text
        for t in doc.texts
        if isinstance(t, TextItem) and t.content_layer == ContentLayer.NOTES
    ]
    assert len(comment_texts) == 4, (
        f"Expected 4 comment texts, got {len(comment_texts)}"
    )

    assert all("[author:" in text for text in comment_texts), (
        "All comments should have author metadata"
    )

    all_text = " ".join(comment_texts)
    assert "John Reviewer (JR)" in all_text, "Expected John Reviewer (JR) in comments"
    assert "Jane Smith (JS)" in all_text, "Expected Jane Smith (JS) in comments"
    assert "sample reviewer comment" in all_text, "Expected original comment text"
    assert "sample response" in all_text, "Expected reply comment text"

    jr_comments = [t for t in comment_texts if "John Reviewer (JR)" in t]
    js_comments = [t for t in comment_texts if "Jane Smith (JS)" in t]
    assert len(jr_comments) == 1, f"Expected 1 comment from JR, got {len(jr_comments)}"
    assert len(js_comments) == 3, f"Expected 3 comments from JS, got {len(js_comments)}"


def test_comments_respect_page_range() -> None:
    """Test that comments are only extracted for slides within page_range."""
    path = Path("./tests/data/pptx/sources/powerpoint_comments.pptx")
    converter = get_converter()

    doc: DoclingDocument = converter.convert(path, page_range=(1, 1)).document

    comment_groups = [g for g in doc.groups if g.name.startswith("comment-")]
    assert len(comment_groups) == 2, (
        f"Expected 2 comment groups from slide 1, got {len(comment_groups)}"
    )

    assert all("slide1" in g.name for g in comment_groups), (
        "Comments should only be from slide 1 when page_range is (1,1)"
    )

    doc3: DoclingDocument = converter.convert(path, page_range=(3, 3)).document

    comment_groups3 = [g for g in doc3.groups if g.name.startswith("comment-")]
    assert len(comment_groups3) == 2, (
        f"Expected 2 comment groups from slide 3, got {len(comment_groups3)}"
    )

    assert all("slide3" in g.name for g in comment_groups3), (
        "Comments should only be from slide 3 when page_range is (3,3)"
    )

    doc2: DoclingDocument = converter.convert(path, page_range=(2, 2)).document
    comment_groups2 = [g for g in doc2.groups if g.name.startswith("comment-")]
    assert len(comment_groups2) == 0, (
        f"Expected 0 comment groups from slide 2, got {len(comment_groups2)}"
    )


def test_pptx_unrecognized_shape_type():
    """PPTX with a <p:sp> that has no geometry should not crash.

    python-pptx raises NotImplementedError from Shape.shape_type for shapes
    that aren't placeholders, autoshapes, textboxes, or freeforms. The
    backend should skip the unrecognized shape gracefully and still extract
    text from the rest of the presentation.

    Ref: https://github.com/docling-project/docling/issues/3308
    """
    converter = get_converter()
    pptx_path = Path("./tests/data/pptx/sources/powerpoint_unrecognized_shape.pptx")

    conv_result: ConversionResult = converter.convert(pptx_path)
    doc: DoclingDocument = conv_result.document

    pred_md = doc.export_to_markdown()

    # Normal slide content should still be extracted
    assert "Q3 Revenue Summary" in pred_md
    assert "Enterprise segment" in pred_md
    assert "Key Metrics" in pred_md
    assert "Next Steps" in pred_md


def test_pptx_malformed_picture_shapes():
    """PPTX with malformed <p:pic> shapes should not crash conversion.

    python-pptx's shape.image accessor raises three distinct exceptions on
    picture shapes that slip past other tools' parsers (Keynote/Google Drive
    open these files fine): InvalidXmlError when <p:blipFill> is missing,
    KeyError when <a:blip r:embed> points at an unknown relationship, and
    AttributeError when the embedded part's content-type isn't an image.

    The backend should skip each malformed picture with a warning and still
    extract text from the slides.
    """
    converter = get_converter()
    pptx_path = Path("./tests/data/pptx/sources/powerpoint_malformed_pictures.pptx")

    with pytest.warns(UserWarning, match="Skipping malformed picture shape"):
        conv_result: ConversionResult = converter.convert(pptx_path)

    doc: DoclingDocument = conv_result.document

    pred_md = doc.export_to_markdown()
    assert "Slide With Missing BlipFill" in pred_md
    assert "Slide With Dangling Rel" in pred_md
    assert "Slide With Wrong Content Type" in pred_md


def test_pptx_left_flush_shape_keeps_own_bbox(tmp_path: Path):
    """A shape positioned at x = 0 EMU must keep its own bounding box.

    shape.left is an Emu, an int subclass, so a left-flush shape made the
    old truthiness check fall into the position-unknown fallback and its
    provenance bbox covered the entire slide.
    """
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    flush = slide.shapes.add_textbox(Inches(0), Inches(1), Inches(3), Inches(1))
    flush.text_frame.text = "flush left"
    pptx_path = tmp_path / "flush_left.pptx"
    prs.save(pptx_path)

    doc = get_converter().convert(pptx_path).document

    item = next(t for t in doc.texts if t.text == "flush left")
    bbox = item.prov[0].bbox
    assert (bbox.l, bbox.r) == (0, Inches(3))
    assert abs(bbox.t - bbox.b) == Inches(1)
    assert bbox.r != prs.slide_width


def test_pptx_page_range():
    converter = get_converter()
    pptx_path = Path("./tests/data/pptx/sources/powerpoint_sample.pptx")

    conv_result: ConversionResult = converter.convert(pptx_path, page_range=(2, 2))

    assert conv_result.input.page_count == 3
    assert conv_result.document.num_pages() == 1
    assert list(conv_result.document.pages.keys()) == [2]

    pred_md = conv_result.document.export_to_markdown()
    assert "Second slide title" in pred_md
    assert "Test Table Slide" not in pred_md
    assert "List item4" not in pred_md


def test_chart_parsed_as_classified_picture_with_data():
    """A native PPTX chart becomes a classified picture carrying its data.

    ``pptx_chart.pptx`` holds two slides:

    * Slide 1 — a 2-D clustered-column chart titled "Wild Duck Observations by
      Year" with two series over four years. It should convert to a PictureItem
      classified as a bar chart, captioned with the chart title, and carrying
      the chart's plotted numbers reconstructed as a table::

          | <blank> | Freshwater Ducks | Saltwater Ducks |
          | 2019    | 120              | 80              |
          ...
          | 2022    | 175              | 130             |

    * Slide 2 — a 3-D bar chart (``c:bar3DChart``) for which python-pptx has no
      registered element class. It should degrade gracefully: the chart is
      emitted as a PictureItem with no tabular data.
    """
    converter = get_converter()
    doc = converter.convert(CHART_PPTX).document

    pictures = list(doc.pictures)
    assert len(pictures) == 2, f"Expected two chart pictures, got {len(pictures)}"

    # --- slide 1: 2-D chart with full data ---
    duck_pic = next(
        p for p in pictures if p.caption_text(doc) == "Wild Duck Observations by Year"
    )
    assert (
        duck_pic.meta.classification.predictions[0].class_name
        == PictureClassificationLabel.BAR_CHART
    )
    chart_data = duck_pic.meta.tabular_chart.chart_data
    assert (chart_data.num_rows, chart_data.num_cols) == (5, 3)
    grid = {
        (cell.start_row_offset_idx, cell.start_col_offset_idx): cell.text
        for cell in chart_data.table_cells
    }
    assert grid[(0, 1)] == "Freshwater Ducks"
    assert grid[(0, 2)] == "Saltwater Ducks"
    assert grid[(1, 0)] == "2019"
    assert grid[(4, 0)] == "2022"
    assert grid[(4, 1)] == "175"
    assert grid[(4, 2)] == "130"

    # --- slide 2: 3-D chart degrades gracefully (no tabular data, no crash) ---
    revenue_pic = next(
        p for p in pictures if p.caption_text(doc) != "Wild Duck Observations by Year"
    )
    assert revenue_pic.meta is not None
    assert revenue_pic.meta.tabular_chart is None


def test_chart_image_not_rendered_by_default():
    """Charts carry classification and data but no image unless opted in.

    render_chart_images defaults to False, so chart pictures keep their
    classification and reconstructed data but no pixels. This guards the promise
    that the feature does not change default output size for existing users.
    """
    converter = get_converter()
    doc = converter.convert(CHART_PPTX).document

    for picture in doc.pictures:
        assert picture.image is None, (
            "chart picture should have no image when render_chart_images is off"
        )


def test_chart_enrichment_skips_image_when_pages_empty():
    """Image enrichment skips native charts without an embedded or page image."""
    format_options = {
        InputFormat.PPTX: PowerpointFormatOption(pipeline_cls=_ChartEnrichmentPipeline)
    }
    converter = DocumentConverter(
        allowed_formats=[InputFormat.PPTX], format_options=format_options
    )

    result = converter.convert(CHART_PPTX, raises_on_error=True)

    pictures = list(result.document.pictures)
    assert len(pictures) == 2
    assert all(p.image is None for p in pictures)


def test_chart_image_rendering(libreoffice_available):
    """render_chart_images=True attaches a LibreOffice-rendered image.

    LibreOffice output is not byte-stable and the cropped image size depends on
    the LibreOffice version, so pixels are not compared against groundtruth. We
    assert the Duck Survey picture (slide 1) gains a non-trivial image while
    keeping the classification and tabular data. Requires LibreOffice; skipped
    when it is not installed.
    """
    if not libreoffice_available:
        pytest.skip("LibreOffice is not installed — chart rendering cannot be tested")

    options = MsPowerpointBackendOptions(render_chart_images=True)
    format_options = {InputFormat.PPTX: PowerpointFormatOption(backend_options=options)}
    converter = DocumentConverter(
        allowed_formats=[InputFormat.PPTX], format_options=format_options
    )
    doc = converter.convert(CHART_PPTX).document

    pictures = list(doc.pictures)
    assert len(pictures) == 2, f"Expected two chart pictures, got {len(pictures)}"

    duck_pic = next(
        p for p in pictures if p.caption_text(doc) == "Wild Duck Observations by Year"
    )
    assert (
        duck_pic.meta.classification.predictions[0].class_name
        == PictureClassificationLabel.BAR_CHART
    )
    assert duck_pic.meta.tabular_chart is not None

    image = duck_pic.get_image(doc=doc)
    assert image is not None, "chart picture should carry a rendered image"
    assert image.width > 50 and image.height > 50, (
        f"rendered chart image is implausibly small: {image.size}"
    )


def test_pptx_shapes_are_sorted_by_visual_position():
    class FakeShape:
        def __init__(self, name, top=None, left=None):
            self.name = name
            self.top = top
            self.left = left

    class BadPositionShape:
        @property
        def top(self):
            raise ValueError("bad position")

    backend = object.__new__(MsPowerpointDocumentBackend)

    same_row_right = FakeShape("same-row-right", top=100, left=300)
    lower_left = FakeShape("lower-left", top=200000, left=100)
    same_row_left = FakeShape("same-row-left", top=1000, left=100)
    unpositioned = FakeShape("unpositioned")

    ordered_shapes = backend._iter_shapes_by_position(
        [lower_left, same_row_right, unpositioned, same_row_left]
    )

    assert [shape.name for shape in ordered_shapes] == [
        "same-row-left",
        "same-row-right",
        "lower-left",
        "unpositioned",
    ]
    assert backend._get_shape_position(BadPositionShape(), "top") is None


def test_pptx_row_grouping_uses_sliding_window():
    """Shapes in a contiguous band should all land in the same row.

    With a fixed-anchor strategy, shapes at tops 0, 40000, and 80000 EMUs
    (each 40000 apart, within the 45720 EMU tolerance) would be split: the
    third shape is 80000 EMUs from the first anchor (0), exceeding tolerance.
    The sliding-window strategy compares each shape against its immediate
    predecessor, so all three end up in the same row and are sorted by left.
    """

    class FakeShape:
        def __init__(self, name, top, left):
            self.name = name
            self.top = top
            self.left = left

    backend = object.__new__(MsPowerpointDocumentBackend)

    # Three shapes in a contiguous band, each 40 000 EMUs apart.
    # Fixed-anchor would split them; sliding-window keeps them together.
    a = FakeShape("a", top=0, left=200)
    b = FakeShape("b", top=40000, left=100)
    c = FakeShape("c", top=80000, left=300)
    # This shape is more than one tolerance step from c, so it forms a new row.
    d = FakeShape("d", top=200000, left=100)

    ordered = [s.name for s in backend._iter_shapes_by_position([d, c, a, b])]

    # a, b, c are in the same row sorted left-to-right; d is in its own row.
    assert ordered == ["b", "a", "c", "d"]
